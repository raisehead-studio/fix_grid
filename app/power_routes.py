from flask import Blueprint, request, jsonify, send_file
from flask_login import login_required, current_user
import sqlite3
import random
from datetime import datetime
from io import BytesIO
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from collections import defaultdict

power_bp = Blueprint('power_bp', __name__)

def get_db():
    conn = sqlite3.connect("kao_power_water.db")
    conn.row_factory = sqlite3.Row
    return conn

@power_bp.route('/api/power_reports', methods=['GET'])
@login_required
def get_power_reports():
    try:
        conn = sqlite3.connect("kao_power_water.db", timeout=10)
        cursor = conn.cursor()
        if current_user.role_id == 4:  # 里幹事
            cursor.execute("""
                SELECT po.*, d.name, v.name
                FROM power_reports po
                LEFT JOIN districts d ON po.district_id = d.id
                LEFT JOIN villages v ON po.village_id = v.id
                WHERE po.deleted_at IS NULL AND po.created_by = ?
                ORDER BY po.created_at
            """, (current_user.id,))
        else:
            cursor.execute("""
                SELECT po.*, d.name, v.name
                FROM power_reports po
                LEFT JOIN districts d ON po.district_id = d.id
                LEFT JOIN villages v ON po.village_id = v.id
                WHERE po.deleted_at IS NULL
                ORDER BY po.created_at
            """)
        rows = cursor.fetchall()
        data = [dict(
            id=row[0],
            district_id=row[1],
            district=row[-2],
            village_id=row[2],
            village=row[-1],
            location=row[3],
            reason=row[4],
            count=row[5],
            contact=row[6],
            phone=row[7],
            report_status=row[12],
            report_restored_at=row[13],
            taipower_status=row[14],
            taipower_description=row[16],  # 這是 taipower_note 欄位
            taipower_eta_hours=row[17],
            taipower_support=row[18],
            taipower_restored_at=row[15],
            created_at=row[9]
        ) for row in rows]
        return jsonify(data)
    except Exception as e:
        return jsonify({'error': str(e)}), 500
    finally:
        if 'conn' in locals():
            conn.close()

@power_bp.route("/api/power_reports", methods=["POST"])
@login_required
def create_power_report():
    data = request.get_json()
    try:
        conn = sqlite3.connect("kao_power_water.db", timeout=10)
        cursor = conn.cursor()
        cursor.execute("""
            INSERT INTO power_reports (district_id, village_id, location, reason, count, contact_name, contact_phone, created_by)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?)
        """, (
            current_user.district_id,
            data["village_id"],
            data["location"],
            data["reason"],
            data["count"],
            data["contact_name"],
            data["contact_phone"],
            current_user.id
        ))
        conn.commit()
        return jsonify({"status": "ok"})
    except Exception as e:
        return jsonify({'error': str(e)}), 500
    finally:
        if 'conn' in locals():
            conn.close()

@power_bp.route('/api/power_reports/<int:id>/update_report', methods=['POST'])
@login_required
def update_power_outage_report(id):
    data = request.get_json()
    try:
        conn = sqlite3.connect("kao_power_water.db", timeout=10)
        cursor = conn.cursor()
        
        # 如果停電戶數為 0，自動將狀態改為已復電
        if data['count'] == 0:
            cursor.execute("""
                UPDATE power_reports
                SET location = ?, reason = ?, count = ?, contact_name = ?, contact_phone = ?, 
                    report_status = 1, report_restored_at = current_timestamp, updated_at = datetime('now')
                WHERE id = ? AND report_status = 0
            """, (
                data['location'],
                data['reason'],
                data['count'],
                data['contact_name'],
                data['contact_phone'],
                id
            ))
        else:
            cursor.execute("""
                UPDATE power_reports
                SET location = ?, reason = ?, count = ?, contact_name = ?, contact_phone = ?, updated_at = datetime('now')
                WHERE id = ? AND report_status = 0
            """, (
                data['location'],
                data['reason'],
                data['count'],
                data['contact_name'],
                data['contact_phone'],
                id
            ))
        conn.commit()
        return jsonify({"status": "ok"})
    except Exception as e:
        return jsonify({'error': str(e)}), 500
    finally:
        if 'conn' in locals():
            conn.close()

@power_bp.route('/api/power_reports/<int:id>/toggle_report_status', methods=['POST'])
@login_required
def toggle_report_status(id):
    try:
        conn = sqlite3.connect("kao_power_water.db", timeout=10)
        cursor = conn.cursor()
        cursor.execute("""
            UPDATE power_reports
            SET report_status = 1, report_restored_at = current_timestamp
            WHERE id = ? AND report_status = 0
        """, (id,))
        conn.commit()
        return jsonify({"status": "restored"})
    except Exception as e:
        return jsonify({'error': str(e)}), 500
    finally:
        if 'conn' in locals():
            conn.close()

@power_bp.route('/api/power_reports/<int:id>/update_taipower', methods=['POST'])
@login_required
def update_taipower_status(id):
    data = request.get_json()
    try:
        conn = sqlite3.connect("kao_power_water.db", timeout=10)
        cursor = conn.cursor()
        cursor.execute("""
            UPDATE power_reports
            SET taipower_note = ?, taipower_eta_hours = ?, taipower_support = ?, taipower_restored_at = current_timestamp
            WHERE id = ?
        """, (
            data['taipower_note'], data['taipower_eta_hours'], data['taipower_support'], id
        ))
        conn.commit()
        return jsonify({"status": "ok"})
    except Exception as e:
        return jsonify({'error': str(e)}), 500
    finally:
        if 'conn' in locals():
            conn.close()

@power_bp.route('/api/power_reports/<int:id>/toggle_taipower_status', methods=['POST'])
@login_required
def toggle_taipower_status(id):
    data = request.get_json()
    if not data or 'taipower_status' not in data:
        return jsonify({"error": "Missing taipower_status"}), 400

    try:
        status = int(data['taipower_status'])
        if status not in (0, 1):
            raise ValueError
    except ValueError:
        return jsonify({"error": "Invalid taipower_status, must be 0 or 1"}), 400

    try:
        conn = sqlite3.connect("kao_power_water.db", timeout=10)
        cursor = conn.cursor()
        cursor.execute("""
            UPDATE power_reports
            SET taipower_status = ?, taipower_restored_at = current_timestamp
            WHERE id = ?
        """, (status, id))
        conn.commit()
        return jsonify({"status": "ok", "taipower_status": status})
    except Exception as e:
        return jsonify({'error': str(e)}), 500
    finally:
        if 'conn' in locals():
            conn.close()

@power_bp.route("/api/power_reports/export-power-report", methods=['POST'])
@login_required
def export_power_report():
    payload = request.get_json()
    data = payload.get("data", [])

    gov_data = []
    tp_data = []

    for row in data:
        if row.get("gov_count", 0) > 0:
            gov_data.append({
                "district": row["district"],
                "village": row["village"],
                "gov_count": row["gov_count"]
            })
        if row.get("tp_count", 0) > 0:
            tp_data.append({
                "district": row["district"],
                "village": row["village"],
                "tp_count": row["tp_count"]
            })

    wb = Workbook()
    ws_district = wb.active
    ws_district.title = "以區分類"
    generate_summary_sheet(ws_district, gov_data, tp_data, by="district")


    ws_village = wb.create_sheet(title="以里分類")
    generate_summary_sheet(ws_village, gov_data, tp_data, by="village")

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    now = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
    filename = f"停電統計_{now}.xlsx"

    return send_file(
        output,
        as_attachment=True,
        download_name=filename,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )

def generate_summary_sheet(ws, gov_data, tp_data, by="village"):
    if by != "district":
        label_col = "里"
        ws.append([label_col, "公所回報", "台電公司"])

        counter = defaultdict(lambda: {"gov": 0, "tp": 0})

        for row in gov_data:
            key = row["village"]
            counter[key]["gov"] += row.get("gov_count", 0)

        for row in tp_data:
            key = row["village"]
            counter[key]["tp"] += row.get("tp_count", 0)

        # 排序依照 max(gov, tp) 由大到小
        all_rows = [{
            "village": k,
            "gov": v["gov"],
            "tp": v["tp"]
        } for k, v in counter.items()]
        all_rows.sort(key=lambda x: max(x["gov"], x["tp"]), reverse=True)

        for row in all_rows:
            ws.append([row["village"], row["gov"], row["tp"]])

        # 圖表
        row_count = len(all_rows) + 1
        data_ref = Reference(ws, min_col=2, max_col=3, min_row=1, max_row=row_count)
        cats_ref = Reference(ws, min_col=1, min_row=2, max_row=row_count)
        chart = BarChart()
        chart.title = "里 停電戶數對比"
        chart.y_axis.title = "戶數"
        chart.x_axis.title = "里"
        chart.add_data(data_ref, titles_from_data=True)
        chart.set_categories(cats_ref)
        chart.width = 28
        chart.height = 15
        ws.add_chart(chart, "F5")
        return

    # --- 以下為 district 模式 ---
    ws.append(["行政區", "公所回報", "戶數", "台電官網資料", "戶數"])

    gov_map = defaultdict(lambda: {"count": 0, "villages": set()})
    tp_map = defaultdict(lambda: {"count": 0, "villages": set()})

    for row in gov_data:
        d = row["district"]
        gov_map[d]["count"] += row.get("gov_count", 0)
        gov_map[d]["villages"].add(row["village"])

    for row in tp_data:
        d = row["district"]
        tp_map[d]["count"] += row.get("tp_count", 0)
        tp_map[d]["villages"].add(row["village"])

    all_districts = sorted(set(gov_map.keys()) | set(tp_map.keys()))

    rows = []
    for d in all_districts:
        gov = gov_map[d]
        tp = tp_map[d]

        gov_count = gov["count"] if d in gov_map else 0
        tp_count = tp["count"] if d in tp_map else 0
        if gov_count == 0 and tp_count == 0:
            continue

        gov_villages = "、".join(sorted(gov["villages"])) if gov["villages"] else ""
        tp_villages = "、".join(sorted(tp["villages"])) if tp["villages"] else ""

        rows.append({
            "district": d,
            "gov_villages": gov_villages,
            "gov_count": gov_count,
            "tp_villages": tp_villages,
            "tp_count": tp_count
        })

    # 排序
    rows.sort(key=lambda r: max(r["gov_count"], r["tp_count"]), reverse=True)

    for r in rows:
        ws.append([
            r["district"],
            r["gov_villages"],
            r["gov_count"],
            r["tp_villages"],
            r["tp_count"]
        ])

    # 加圖表（用戶數）
    row_count = len(rows) + 1
    data_ref = Reference(ws, min_col=3, max_col=5, min_row=1, max_row=row_count)
    cats_ref = Reference(ws, min_col=1, min_row=2, max_row=row_count)
    chart = BarChart()
    chart.title = "行政區 停電戶數對比"
    chart.y_axis.title = "戶數"
    chart.x_axis.title = "行政區"
    chart.add_data(data_ref, titles_from_data=True)
    chart.set_categories(cats_ref)
    chart.width = 28
    chart.height = 15
    ws.add_chart(chart, "G5")

@power_bp.route('/api/power_stats', methods=['GET'])
@login_required
def get_power_stats():
    try:
        conn = get_db()
        cursor = conn.cursor()

        if current_user.role_id == 4:
            cursor.execute("""
                SELECT po.*, d.name, v.name
                FROM power_reports po
                LEFT JOIN districts d ON po.district_id = d.id
                LEFT JOIN villages v ON po.village_id = v.id
                WHERE po.deleted_at IS NULL AND po.created_by = ?
                ORDER BY po.created_at
            """, (current_user.id,))
        else:
            cursor.execute("""
                SELECT po.*, d.name, v.name
                FROM power_reports po
                LEFT JOIN districts d ON po.district_id = d.id
                LEFT JOIN villages v ON po.village_id = v.id
                WHERE po.deleted_at IS NULL
                ORDER BY po.created_at
            """)
        gov_rows = cursor.fetchall()

        gov_data = [dict(
            district_id=row[1],
            village_id=row[2],
            district=row[-2],
            village=row[-1],
            gov_count=row[5]
        ) for row in gov_rows]

        cursor.execute("""
            SELECT tr.district_id, tr.village_id, d.name, v.name, SUM(tr.count)
            FROM taipower_reports tr
            LEFT JOIN districts d ON tr.district_id = d.id
            LEFT JOIN villages v ON tr.village_id = v.id
            GROUP BY tr.district_id, tr.village_id
        """)
        tp_rows = cursor.fetchall()
    except Exception as e:
        return jsonify({'error': str(e)}), 500
    finally:
        if 'conn' in locals():
            conn.close()

    tp_map = {(row[0], row[1]): {'district': row[2], 'village': row[3], 'tp_count': row[4]} for row in tp_rows}

    combined_map = {}
    for row in gov_data:
        key = (row['district_id'], row['village_id'])
        combined_map[key] = {
            'district_id': row['district_id'],
            'village_id': row['village_id'],
            'district': row['district'],
            'village': row['village'],
            'gov_count': row['gov_count'],
            'tp_count': 0
        }

    for key, tp in tp_map.items():
        if key not in combined_map:
            combined_map[key] = {
                'district_id': key[0],
                'village_id': key[1],
                'district': tp['district'],
                'village': tp['village'],
                'gov_count': 0,
                'tp_count': tp['tp_count']
            }
        else:
            combined_map[key]['tp_count'] = tp['tp_count']

    return jsonify(list(combined_map.values()))

@power_bp.route("/api/taipower_reports", methods=["POST"])
@login_required
def create_taipower_report():
    data = request.get_json()
    try:
        conn = get_db()
        cursor = conn.cursor()
        cursor.execute("""
            INSERT INTO taipower_reports (district_id, village_id, count, created_by)
            VALUES (?, ?, ?, ?)
        """, (data["district_id"], data["village_id"], data["count"], current_user.id))
        conn.commit()
        return jsonify({"status": "ok"})
    except Exception as e:
        return jsonify({"error": str(e)}), 500
    finally:
        if 'conn' in locals():
            conn.close()
